# -*- coding: utf-8 -*-
"""tubestudy 提案（第2版・16:9・9枚＋動画）。python-pptx で組む。
大きな文字だけ。1枚に3〜5行。添え書きは置かない。発表者ノートに読み原稿。

  python3 build_deck2.py [出力.pptx]
"""
import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
VIDEO = "/private/tmp/claude-501/-Users-aki-Library-Mobile-Documents-com-apple-CloudDocs-myenglish/4055e3c9-2bd7-4fb6-9add-6e3418902aef/scratchpad/promo/tubestudy-v4.mp4"
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "tubestudy-提案.pptx")

FONT = "Hiragino Sans"
NAVY = RGBColor(0x0F, 0x3F, 0x73)
BLUE = RGBColor(0x00, 0x4E, 0x93)
YEL = RGBColor(0xF9, 0xB9, 0x00)
INK = RGBColor(0x17, 0x15, 0x0F)
INK2 = RGBColor(0x5B, 0x5A, 0x55)
INK3 = RGBColor(0x8A, 0x89, 0x84)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
TINT2 = RGBColor(0xE8, 0xEE, 0xF7)
YTINT = RGBColor(0xFD, 0xF1, 0xC7)
GTINT = RGBColor(0xDF, 0xF1, 0xE4)
LINE = RGBColor(0xD9, 0xDD, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xCA, 0xDC, 0xFC)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

NO_START = "、。」）】,.!?！？ー"
NO_END = "「（【"


def cap(w, pt):
    return max(4, int((w - 0.2) * 72 / pt * 0.95))


def jp(text, n):
    out = []
    for para_ in text.split("\n"):
        cur = ""; i = 0
        while i < len(para_):
            cur += para_[i]; i += 1
            if len(cur) >= n:
                while i < len(para_) and para_[i] in NO_START:
                    cur += para_[i]; i += 1
                j = -1
                for k in range(len(cur) - 1, int(n * 0.45), -1):
                    if cur[k - 1] in "、。」）" and k < len(cur):
                        j = k; break
                if j <= 0:
                    j = len(cur)
                    if cur[j - 1] in NO_END: j -= 1
                out.append(cur[:j]); cur = cur[j:]
        if cur: out.append(cur)
    return out


def set_font(run, size, bold=False, color=INK, name=FONT):
    f = run.font
    f.name = name; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None: el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15, margin=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False; p.alignment = align; p.line_spacing = spacing
        for r in (line if isinstance(line, list) else [line]):
            if isinstance(r, str): r = (r, 18, False, INK)
            s, size, bold, color = (list(r) + [None] * 4)[:4]
            run = p.add_run(); run.text = s
            set_font(run, size or 18, bool(bold), color or INK)
    return tb


def para(slide, x, y, w, s, size, color=INK2, bold=False, spacing=1.35, align=PP_ALIGN.LEFT):
    lines = jp(s, cap(w, size))
    h = len(lines) * size / 72.0 * spacing + 0.08
    text(slide, x, y, w, h, [[(l, size, bold, color)] for l in lines], align=align, spacing=spacing)
    return h


def rect(slide, x, y, w, h, fill, line=None, radius=None, shape=None):
    sp = slide.shapes.add_shape(shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    if radius is not None and sp.adjustments: sp.adjustments[0] = radius
    sp.shadow.inherit = False; sp.text_frame.text = ""
    return sp


def round_pic(pic, adj=0.04):
    spPr = pic._element.spPr
    geom = spPr.find(qn("a:prstGeom")); geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None: av = etree.SubElement(geom, qn("a:avLst"))
    for c in list(av): av.remove(c)
    g = etree.SubElement(av, qn("a:gd")); g.set("name", "adj"); g.set("fmla", "val %d" % int(adj * 50000))


def screen(slide, path, x, y, w, crop_bottom=0.0, border=True, adj=0.03, crop_top=0.0):
    pic = slide.shapes.add_picture(os.path.join(IMG, path), Inches(x), Inches(y), width=Inches(w))
    if crop_bottom or crop_top:
        pic.crop_bottom = crop_bottom; pic.crop_top = crop_top
        pic.height = int(pic.height * (1 - crop_bottom - crop_top))
    round_pic(pic, adj)
    if border: pic.line.color.rgb = LINE; pic.line.width = Pt(0.75)
    return pic


def phone(slide, path, x, y, w, crop_bottom=0.0, crop_top=0.0):
    pad = w * 0.035
    pic = slide.shapes.add_picture(os.path.join(IMG, path), Inches(x + pad), Inches(y + pad), width=Inches(w - 2 * pad))
    if crop_bottom or crop_top:
        pic.crop_bottom = crop_bottom; pic.crop_top = crop_top
        pic.height = int(pic.height * (1 - crop_bottom - crop_top))
    h = Emu(pic.height).inches + 2 * pad
    bezel = rect(slide, x, y, w, h, RGBColor(0x11, 0x11, 0x11), radius=0.12)
    sp = bezel._element; sp.getparent().remove(sp); pic._element.addprevious(sp)
    round_pic(pic, 0.09)
    return h


def mark(slide, x, y, size):
    box = rect(slide, x, y, size, size, BLUE, radius=0.24)
    rect(slide, x + size * 0.19, y + size * 0.25, size * 0.62, size * 0.078, YEL, radius=0.5)
    ring_d = size * 0.44
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + (size - ring_d) / 2), Inches(y + size * 0.53 - ring_d / 2 + size * 0.02), Inches(ring_d), Inches(ring_d))
    o.fill.background(); o.line.color.rgb = WHITE; o.line.width = Pt(size * 5.2); o.shadow.inherit = False
    return box


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def page(slide, n):
    text(slide, 12.3, 7.05, 0.8, 0.3, [[(str(n), 10, False, INK3)]], align=PP_ALIGN.RIGHT)


def title(slide, s, y=0.6, size=40, color=INK):
    text(slide, 0.8, y, 11.8, 1.0, [[(s, size, True, color)]], spacing=1.1)


def num_circle(slide, x, y, d, n, fill=BLUE):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); set_font(r, int(d * 24), True, WHITE)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def pill(slide, x, y, w, h, s, size=16, fill=TINT2, color=BLUE):
    rect(slide, x, y, w, h, fill, radius=0.5)
    text(slide, x, y, w, h, [[(s, size, True, color)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x, y, w, h, size=22):
    text(slide, x, y, w, h, [[("→", size, True, INK3)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def big_rows(slide, rows, x=0.8, y=1.9, w=11.6, size=30, gap=0.42, who=None):
    """大きな文字の行を、番号つきで縦に。who があれば番号の代わりに「生徒」などの札"""
    for i, s in enumerate(rows):
        if who:
            label, fill, col = who[i]
            pill(slide, x, y + 0.02, 1.35, 0.6, label, 20, fill, col)
            tx = x + 1.6
        else:
            num_circle(slide, x, y + 0.04, 0.6, i + 1)
            tx = x + 0.9
        h = para(slide, tx, y, w - (tx - x), s, size, INK, bold=True, spacing=1.25)
        y += max(h, 0.7) + gap
    return y


n = 0


def new(bgc=PAPER):
    global n
    n += 1
    s = prs.slides.add_slide(BLANK); bg(s, bgc)
    if bgc == PAPER: page(s, n)
    return s


# ============================ 1 表紙 ============================
s = new(NAVY)
mark(s, 0.9, 1.5, 1.35)
text(s, 2.5, 1.5, 8, 1.4, [[("tube", 60, True, WHITE), ("study", 60, True, YEL)]], anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
text(s, 0.9, 3.4, 11.5, 1.2, [[("ロンドン校での、試験導入のお願い", 40, True, WHITE)]], spacing=1.1)
text(s, 0.9, 5.7, 11, 0.6, [[("ロンドン校 講師　秋 建太郎", 22, False, ICE)]])
notes(s, "本日は、こういう機会をいただけて本当にありがとうございます。ロンドンに着いたばかりで、まだ日々教わることばかりです。今日は、役に立ちたくて作ったものを見ていただき、試しに使わせていただけないか、というお願いです。10分ほどお時間をください。")

# ============================ 2 現場でやってみて、感じたこと ============================
s = new()
title(s, "現場でやってみて、感じたこと")
big_rows(s, [
    "作文や漢字テストは、集めて、丸をつけて、返す。",
    "オンラインの生徒には、印刷して、赤を入れて、スキャンして、メール。",
    "単元ごとの理解は、先生の記憶が支えている。",
], y=2.0, size=28, gap=0.5)
notes(s, "まず、この数週間、ロンドン校でやってみて感じたことです。先生方が一人ひとりに丁寧に向き合っていることを、日々感じています。同時に、その丁寧さの多くが手作業で支えられていることも感じました。作文や漢字テストは、集めて、丸をつけて、返す。オンラインの生徒には、印刷して、赤を入れて、スキャンして、メールで送る。単元ごとに誰がどこでつまずいているかは、特訓選抜テストのときには分かりますが、日々は先生の記憶に頼っています。私自身がやってみて、この時間を授業の準備に回せたら、と思いました。")

# ============================ 3 ご提案 ============================
s = new()
text(s, 0.8, 0.9, 11.8, 1.2, [[("宿題と成績を、仕組みで後押しする。", 44, True, INK)]], spacing=1.1)
text(s, 0.82, 2.05, 11.8, 0.8, [[("講師の時間をもっと、授業へ充てるためのご提案です。", 28, False, INK2)]], spacing=1.2)
# 3つの端末（先生 iPad／生徒 iPhone／保護者 iPhone）
lx = 0.9
text(s, lx, 3.35, 5.8, 0.5, [[("先生", 22, True, BLUE)]])
screen(s, "t-home.png", lx, 3.85, 5.8, crop_bottom=0.12)
px = 7.4
text(s, px, 3.35, 2.0, 0.5, [[("生徒", 22, True, BLUE)]])
phone(s, "s-home.png", px, 3.85, 1.75, crop_bottom=0.24)
px2 = 10.0
text(s, px2, 3.35, 2.0, 0.5, [[("保護者", 22, True, BLUE)]])
phone(s, "s-parent.png", px2, 3.85, 1.75, crop_bottom=0.24)
notes(s, "そこで、宿題と成績を、仕組みで後押しする。講師の時間をもっと、授業へ充てるためのご提案です。先生はパソコンかiPad、生徒と保護者はスマホで使う、ひとつのアプリです。紙のクラスと混ざったまま使えます。言葉で説明するより、1分ほどの動画を見ていただくのが早いと思います。")

# ============================ 4 動画 ============================
s = new()
poster = os.path.join(IMG, "video-poster.png")
if os.path.exists(VIDEO):
    s.shapes.add_movie(VIDEO, Inches(0.667), Inches(0.375), Inches(12.0), Inches(6.75), poster_frame_image=poster, mime_type="video/mp4")
else:
    screen(s, "video-poster.png", 0.667, 0.375, 12.0, border=False)
notes(s, "（動画 67秒。再生が終わったら 2〜3秒黙って、反応を見る）")

# ============================ 5 実演 ============================
s = new()
title(s, "実演")
big_rows(s, [
    "写真で、出す。",
    "数字を押す。赤を入れて、返す。",
    "評価と、続けた記録が届く。",
    "理解度が、データに残る。",
    "専用ページで、成績を見る。",
], x=0.8, y=1.75, w=8.0, size=26, gap=0.3,
    who=[("生徒", YTINT, INK), ("先生", TINT2, BLUE), ("生徒", YTINT, INK), ("先生", TINT2, BLUE), ("保護者", GTINT, RGBColor(0x1E, 0x6B, 0x3A))])
screen(s, "t-check-b.png", 8.7, 1.85, 4.2, crop_bottom=0.1)
phone(s, "s-fb.png", 11.2, 4.3, 1.5, crop_bottom=0.3)
notes(s, "ここからは、実際の画面で。5つだけです。①生徒が写真で出す。②先生が数字を押し、気になる子には赤を入れて返す。③生徒に評価と、続けた記録が届く。④先生の側には理解度がデータとして残る。⑤保護者は専用ページで成績を見る。順番にお見せします。（ブラウザへ切り替え。手順は「実演の手順」の紙のとおり）")

# ============================ 6 紙の宿題は、そのまま ============================
s = new()
title(s, "紙の宿題は、そのまま。")
big_rows(s, [
    "生徒は、ノートを写真で出すだけ。",
    "オンラインの生徒にも、同じ手順で返せる。",
    "個人情報は、使わない。保存しない。",
], x=0.8, y=1.85, w=11.6, size=30, gap=0.4)
screen(s, "record-card.png", 9.5, 4.75, 3.2, border=False, adj=0.06)
notes(s, "実演のあとに、3つだけ補足します。1つ目、紙はそのままです。生徒は今まで通りノートや原稿用紙に書き、写真で出すだけ。先生は今まで通り目で見て、赤を入れる。変わるのは、集める・配る・記録する、の部分だけです。2つ目、オンラインの生徒にも同じ手順で返せます。印刷・スキャン・メールが無くなります。保護者も同じ画面で、返ってきた答案と評価を見られます。3つ目、個人情報を持ちません。サーバにあるのは番号と学びの記録だけ。名前は塾の手元の紙の名簿にしかありません。本部のシステムにはつながず、教材データも置きません。")

# ============================ 7 お願い ============================
s = new()
title(s, "ロンドン校で、私のクラスから。")
y = big_rows(s, [
    "紙のクラスと、混ざったままで。",
    "いつでも、紙に戻せます。",
    "費用も、契約も、ありません。",
], x=0.8, y=1.85, w=11.6, size=30, gap=0.36)
# 3段階
py = 5.55
pill(s, 0.8, py, 3.2, 0.7, "① 私のクラス", 20)
arrow(s, 4.05, py, 0.5, 0.7)
pill(s, 4.6, py, 3.7, 0.7, "② 使いたい先生のクラス", 20)
arrow(s, 8.35, py, 0.5, 0.7)
pill(s, 8.9, py, 3.6, 0.7, "③ 校舎として判断", 20)
text(s, 0.82, 6.45, 11.6, 0.5, [[("3か月後に、数字でご報告します。", 22, True, INK2)]])
notes(s, "お願いは一つです。ロンドン校で、まず私の担当クラスから、試験的に使わせてください。費用はかかりません。契約の話でもありません。一講師の取り組みとして、校舎の許可をいただきたい。段階は3つ。私のクラス、次に使ってみたい先生のクラス、そして校舎として続けるかの判断。紙のクラスと混ざったままで動きます。合わなければ、いつでも紙に戻せますし、記録は消せます。お願いしたいことは2つ。生徒と保護者へのご案内を、校舎の名前で出させていただくこと。それと、もし可能であれば、単元確認テストの結果を取り込む許可。無くても宿題だけで動きますが、あると一人ひとりの理解度が宿題とつながります。3か月後に、宿題の提出率、チェックが行き届いた割合、返却までの日数、保護者の反応を、数字でご報告します。")

# ============================ 8 締め ============================
s = new(NAVY)
text(s, 0.9, 2.6, 11.6, 1.4, [[("先生の時間を、授業へ。", 48, True, WHITE), ("もっと。", 48, True, YEL)]], spacing=1.1)
text(s, 0.9, 4.3, 11, 0.6, [[("ありがとうございました。", 24, False, ICE)]])
mark(s, 11.2, 5.4, 1.1)
notes(s, "成績が上がり、講師の質が上がり、保護者の評価が上がる。その循環を、人の力に加えて仕組みで支えたい。私はまだ着いたばかりですが、ロンドン校の役に立ちたいという気持ちで作りました。ぜひ、試させてください。ありがとうございました。")

# ============================ 9 ご参考（質疑用） ============================
s = new()
title(s, "ご参考：データの置き場所")
big_rows(s, [
    "サーバにあるのは、番号と記録だけ。",
    "氏名・住所・電話は、ありません。",
    "本部のシステムには、つなぎません。",
    "教材データは、置きません。",
], x=0.8, y=1.85, w=11.6, size=28, gap=0.34)
notes(s, "（質疑用）サーバにあるのは、番号と学びの記録だけです。氏名・住所・電話は入っていません。名前は塾の手元の紙の名簿にしかなく、サーバの中身だけでは誰の記録か分かりません。パスワードも合言葉も元に戻せない形で置き、判定はすべてサーバ側です。自分で攻撃して塞いだ記録もあります。本部のシステムには一切つながず、教材データも置きません。サーバは英国ロンドンにあり、私が個人で用意しています。")

prs.save(OUT)
print("saved", OUT, "slides", n)
