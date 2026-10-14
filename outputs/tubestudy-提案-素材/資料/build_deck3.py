# -*- coding: utf-8 -*-
"""tubestudy 提案（第3版・旧版の様式・10枚）。python-pptx で組む。Google スライドに取り込んで直せる形。

体裁の決まり（旧版と同じ）
  ・文は自分で行に切る（jp）。箱の端まで詰めず、句読点のあとで折る
  ・箇条書きは1つずつ別の箱に置く（Google スライドで重ならないように）
  ・字体は Noto Sans JP（Google スライドにある）
  python3 build_deck3.py [出力.pptx]
"""
import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
VIDEO = "/private/tmp/claude-501/-Users-aki-Library-Mobile-Documents-com-apple-CloudDocs-myenglish/4055e3c9-2bd7-4fb6-9add-6e3418902aef/scratchpad/promo/tubestudy-v4.mp4"
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "tubestudy-提案.pptx")

FONT = "Noto Sans JP"
NAVY = RGBColor(0x0F, 0x3F, 0x73)
BLUE = RGBColor(0x00, 0x4E, 0x93)
BLUE2 = RGBColor(0x0A, 0x6F, 0xC7)
YEL = RGBColor(0xF9, 0xB9, 0x00)
INK = RGBColor(0x17, 0x15, 0x0F)
INK2 = RGBColor(0x5B, 0x5A, 0x55)
INK3 = RGBColor(0x8A, 0x89, 0x84)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
TINT = RGBColor(0xF2, 0xF4, 0xF8)
TINT2 = RGBColor(0xE8, 0xEE, 0xF7)
YTINT = RGBColor(0xFD, 0xF1, 0xC7)
GTINT = RGBColor(0xDF, 0xF1, 0xE4)
LINE = RGBColor(0xD9, 0xDD, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
GREEN = RGBColor(0x1E, 0x6B, 0x3A)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

NO_START = "、。」）】,.!?！？ー"
NO_END = "「（【"


def cap(w, pt):
    return max(4, int((w - 0.2) * 72 / pt * 0.95))


def jp(text, n):
    out = []
    for para_ in text.split("\n"):
        cur = ""; i = 0
        while i < len(para_):
            cur += para_[i]; i += 1
            if len(cur) >= n:
                while i < len(para_) and para_[i] in NO_START:
                    cur += para_[i]; i += 1
                j = -1
                for k in range(len(cur) - 1, int(n * 0.45), -1):
                    if cur[k - 1] in "、。」）" and k < len(cur):
                        j = k; break
                if j <= 0:
                    j = len(cur)
                    if cur[j - 1] in NO_END: j -= 1
                out.append(cur[:j]); cur = cur[j:]
        if cur: out.append(cur)
    return out


def set_font(run, size, bold=False, color=INK, name=FONT):
    f = run.font
    f.name = name; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None: el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15, margin=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False; p.alignment = align; p.line_spacing = spacing
        for r in (line if isinstance(line, list) else [line]):
            if isinstance(r, str): r = (r, 18, False, INK)
            s, size, bold, color = (list(r) + [None] * 4)[:4]
            run = p.add_run(); run.text = s
            set_font(run, size or 18, bool(bold), color or INK)
    return tb


def para(slide, x, y, w, s, size, color=INK2, bold=False, spacing=1.35, align=PP_ALIGN.LEFT):
    lines = jp(s, cap(w, size))
    h = len(lines) * size / 72.0 * spacing + 0.08
    text(slide, x, y, w, h, [[(l, size, bold, color)] for l in lines], align=align, spacing=spacing)
    return h


def rect(slide, x, y, w, h, fill, line=None, radius=None, shape=None):
    sp = slide.shapes.add_shape(shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE),
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(0.75)
    if radius is not None and sp.adjustments: sp.adjustments[0] = radius
    sp.shadow.inherit = False; sp.text_frame.text = ""
    return sp


def round_pic(pic, adj=0.04):
    spPr = pic._element.spPr
    geom = spPr.find(qn("a:prstGeom")); geom.set("prst", "roundRect")
    av = geom.find(qn("a:avLst"))
    if av is None: av = etree.SubElement(geom, qn("a:avLst"))
    for c in list(av): av.remove(c)
    g = etree.SubElement(av, qn("a:gd")); g.set("name", "adj"); g.set("fmla", "val %d" % int(adj * 50000))


def screen(slide, path, x, y, w, crop_bottom=0.0, border=True, adj=0.03, crop_top=0.0):
    pic = slide.shapes.add_picture(os.path.join(IMG, path), Inches(x), Inches(y), width=Inches(w))
    if crop_bottom or crop_top:
        pic.crop_bottom = crop_bottom; pic.crop_top = crop_top
        pic.height = int(pic.height * (1 - crop_bottom - crop_top))
    round_pic(pic, adj)
    if border: pic.line.color.rgb = LINE; pic.line.width = Pt(0.75)
    return pic


def phone(slide, path, x, y, w, crop_bottom=0.0, crop_top=0.0):
    pad = w * 0.035
    pic = slide.shapes.add_picture(os.path.join(IMG, path), Inches(x + pad), Inches(y + pad), width=Inches(w - 2 * pad))
    if crop_bottom or crop_top:
        pic.crop_bottom = crop_bottom; pic.crop_top = crop_top
        pic.height = int(pic.height * (1 - crop_bottom - crop_top))
    h = Emu(pic.height).inches + 2 * pad
    bezel = rect(slide, x, y, w, h, RGBColor(0x11, 0x11, 0x11), radius=0.12)
    sp = bezel._element; sp.getparent().remove(sp); pic._element.addprevious(sp)
    round_pic(pic, 0.09)
    return h


def mark(slide, x, y, size):
    box = rect(slide, x, y, size, size, BLUE, radius=0.24)
    rect(slide, x + size * 0.19, y + size * 0.25, size * 0.62, size * 0.078, YEL, radius=0.5)
    ring_d = size * 0.44
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + (size - ring_d) / 2), Inches(y + size * 0.53 - ring_d / 2 + size * 0.02), Inches(ring_d), Inches(ring_d))
    o.fill.background(); o.line.color.rgb = WHITE; o.line.width = Pt(size * 5.2); o.shadow.inherit = False
    return box


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def page(slide, n):
    text(slide, 12.3, 7.05, 0.8, 0.3, [[(str(n), 10, False, INK3)]], align=PP_ALIGN.RIGHT)


def title(slide, s, sub=None, y=0.55):
    text(slide, 0.7, y, 12, 0.9, [[(s, 34, True, INK)]], spacing=1.1)
    if sub:
        text(slide, 0.72, y + 0.82, 11.8, 0.5, [[(sub, 15, False, INK2)]])


def num_circle(slide, x, y, d, n, fill=BLUE):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); set_font(r, int(d * 22), True, WHITE)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def pill(slide, x, y, w, h, s, size=14, fill=TINT2, color=BLUE):
    rect(slide, x, y, w, h, fill, radius=0.5)
    text(slide, x, y, w, h, [[(s, size, True, color)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def bullets(slide, x, y, w, rows, size=14, gap=0.14, color=INK):
    """短い行を1つずつ別の箱に。点は青い小さな丸"""
    for s in rows:
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y + size / 72.0 * 0.62 - 0.06), Inches(0.12), Inches(0.12))
        c.fill.solid(); c.fill.fore_color.rgb = BLUE; c.line.fill.background(); c.shadow.inherit = False
        h = para(slide, x + 0.28, y, w - 0.28, s, size, color, spacing=1.35)
        y += h + gap
    return y


def block3(slide, y, cards, h, head_size=20):
    """講師・生徒・保護者の3つの箱。cards = [(見出し, 札の色, 文字色, [行…], 画像の置き方 or None)]"""
    for i, (hd, fill, col, rows, pic) in enumerate(cards):
        x = 0.7 + i * 4.05
        rect(slide, x, y, 3.75, h, TINT, radius=0.08)
        pill(slide, x + 0.3, y + 0.3, 1.2, 0.42, hd, 15, fill, col)
        ty = y + 0.95
        if pic:
            kind, path, w_, crop = pic
            if kind == "screen":
                screen(slide, path, x + 0.3, ty, w_, crop_bottom=crop)
                ty += w_ * 900 / 1440 * (1 - crop) + 0.3
            else:
                ph = phone(slide, path, x + 0.3, ty, w_, crop_bottom=crop)
                ty += ph + 0.3
        bullets(slide, x + 0.3, ty, 3.3, rows, 14)


n = 0


def new(bgc=PAPER):
    global n
    n += 1
    s = prs.slides.add_slide(BLANK); bg(s, bgc)
    if bgc == PAPER: page(s, n)
    return s


# ============================ 1 表紙 ============================
s = new(NAVY)
mark(s, 0.9, 1.35, 1.35)
text(s, 2.55, 1.42, 6, 1.2, [[("tube", 44, True, WHITE), ("study", 44, True, YEL)]], spacing=1.0)
text(s, 0.9, 3.15, 11.5, 1.6, [[("宿題と成績を、仕組みで後押しする。", 44, True, WHITE)]], spacing=1.15)
text(s, 0.92, 4.35, 11, 0.6, [[("講師の時間をもっと、授業へ充てるためのご提案", 18, False, ICE)]])
text(s, 0.92, 6.35, 8, 0.5, [[("早稲田アカデミー ロンドン校　　秋 建太郎　　2026年9月", 12, False, RGBColor(0xA9, 0xBE, 0xDD))]])
notes(s, "本日は、こういう機会をいただけて本当にありがとうございます。ロンドンに着いたばかりで、日々教わることばかりです。今日は、役に立ちたくて作ったものを見ていただき、ロンドン校で試しに使わせていただけないか、というお願いです。10分ほどお時間をください。")

# ============================ 2 「授業準備」以外の大切な仕事 ============================
s = new()
title(s, "「授業準備」以外の大切な仕事")
cards = [("宿題管理", "集めて、見て、返す。\n毎回、全員に。"),
         ("成績管理", "得意と苦手を、\n生徒ごとに。"),
         ("保護者対応", "提出と伸びを、\n家庭に伝える。")]
for i, (h, d) in enumerate(cards):
    x = 0.7 + i * 4.05
    rect(s, x, 1.9, 3.75, 3.4, TINT, radius=0.08)
    num_circle(s, x + 0.35, 2.25, 0.55, i + 1)
    text(s, x + 0.35, 3.0, 3.1, 0.6, [[(h, 24, True, INK)]])
    para(s, x + 0.35, 3.7, 3.1, d, 16, INK2, spacing=1.45)
text(s, 0.7, 5.75, 12, 0.5, [[("生徒が増えるほど、この時間も増える。", 22, True, BLUE)]])
text(s, 0.72, 6.3, 12, 0.5, [[("→ 質を保ったまま、効率的にしていきたい。", 18, False, INK2)]])
notes(s, "授業の準備のほかにも、大切な仕事があります。宿題を集めて、見て、返す。得意と苦手を生徒ごとに見る。提出と伸びを家庭に伝える。どれも一人ひとりに向き合う仕事で、生徒が増えるほど、この時間も増えます。質を保ったまま、効率的にしていきたい。それが出発点です。")

# ============================ 3 3つの目的 ============================
s = new()
title(s, "3つの目的")
goals = [("講師の時間を、授業へ", "宿題チェック・成績管理の負担を軽くし、\nいちばん大切な授業にリソースを投下できるように。"),
         ("一人ひとりに、本当に向き合う「仕組み」", "得意・苦手・提出・先生のことばが生徒ごとに残り、\n誰が担当しても同じ質で向き合える。"),
         ("マンパワーで達成した生徒数を、「仕組み」で支える", "誰もが使える、宿題・成績管理の仕組みをつくる。")]
for i, (h, d) in enumerate(goals):
    y = 1.85 + i * 1.55
    num_circle(s, 0.75, y + 0.05, 0.62, i + 1, YEL if i == 1 else BLUE)
    text(s, 1.6, y - 0.02, 7.8, 0.6, [[(h, 23, True, INK)]])
    para(s, 1.62, y + 0.55, 7.6, d, 14, INK2)
rect(s, 9.55, 1.7, 3.1, 4.9, TINT, radius=0.08)
mark(s, 10.35, 2.35, 1.5)
text(s, 9.55, 4.15, 3.1, 0.6, [[("tube", 26, True, BLUE), ("study", 26, True, YEL)]], align=PP_ALIGN.CENTER, spacing=1.0)
para(s, 9.8, 4.85, 2.6, "講師はパソコン、\n生徒と保護者はスマホ。\n1つのアプリで、\n紙のクラスと併走できる。", 12, INK2, align=PP_ALIGN.CENTER)
notes(s, "目的は3つです。講師の時間を授業へ。一人ひとりに本当に向き合う仕組み。そして、これまで人の力で達成してきた生徒数を、仕組みで支えること。誰もが使える、宿題と成績管理の仕組みをつくりたいと考えています。")

# ============================ 4 tubestudy とは ============================
s = new()
mark(s, 0.9, 0.9, 1.1)
text(s, 2.2, 0.95, 6, 1.0, [[("tube", 40, True, BLUE), ("study", 40, True, YEL)]], spacing=1.0)
text(s, 0.9, 2.35, 11.8, 1.6, [[("生徒、講師、保護者が毎日使える、", 34, True, INK)], [("宿題・成績管理サービス。", 34, True, INK)]], spacing=1.25)
lx = 0.9
text(s, lx, 4.35, 5.8, 0.4, [[("講師", 15, True, BLUE)]])
screen(s, "t-home.png", lx, 4.75, 5.6, crop_bottom=0.42)
text(s, 7.2, 4.35, 2.0, 0.4, [[("生徒", 15, True, BLUE)]])
phone(s, "s-home.png", 7.2, 4.75, 1.7, crop_bottom=0.45)
text(s, 9.75, 4.35, 2.0, 0.4, [[("保護者", 15, True, BLUE)]])
phone(s, "s-parent.png", 9.75, 4.75, 1.7, crop_bottom=0.45)
notes(s, "そこで作ったのが tubestudy です。生徒、講師、保護者が毎日使える、宿題と成績管理のサービスです。講師はパソコンか iPad、生徒と保護者はスマホ。ひとつのアプリで、紙のクラスと混ざったまま使えます。言葉で説明するより、1分ほどの動画を見ていただくのが早いと思います。")

# ============================ 5 動画 ============================
s = new()
poster = os.path.join(IMG, "video-poster.png")
if os.path.exists(VIDEO):
    s.shapes.add_movie(VIDEO, Inches(0.667), Inches(0.375), Inches(12.0), Inches(6.75), poster_frame_image=poster, mime_type="video/mp4")
else:
    screen(s, "video-poster.png", 0.667, 0.375, 12.0, border=False)
notes(s, "（動画 67秒。Google スライドでは Drive の tubestudy-紹介動画.mp4 を挿入する。終わったら 2〜3秒黙って、反応を見る）")

# ============================ 6 できること ============================
s = new()
title(s, "tubestudy でできること", "このあと、実際の画面でお見せします")
block3(s, 1.85, [
    ("講師", TINT2, BLUE, ["見るべき宿題が、上に並ぶ", "写真を見て、1〜5 を押す", "iPad で赤を入れて、返す", "理解度から、宿題を足す"], ("screen", "t-check-b.png", 3.15, 0.45)),
    ("生徒", YTINT, INK, ["今日やる宿題が、ひと目で", "ノートを撮って、出す", "評価と、赤の入った答案が届く", "毎日テストで、苦手をつぶす"], ("phone", "s-home.png", 1.35, 0.55)),
    ("保護者", GTINT, GREEN, ["成績と、先生のことばを見る", "相談を、先生へ送れる", "返事も、同じ画面に"], ("phone", "s-parent.png", 1.35, 0.55)),
], 5.1)
notes(s, "できることを、講師・生徒・保護者に分けて。講師は、見るべき宿題が今日やることに並び、写真を見て1〜5を押す。気になる子には iPad で赤を入れて返す。理解度から宿題を足せる。生徒は、今日やる宿題がひと目で分かり、ノートを撮って出すだけ。評価と赤の入った答案が届き、毎日テストで苦手をつぶす。保護者は、成績・提出・先生のことばをいつでも見られ、相談を送れて、返事も同じ画面に届きます。ここからは、実際の画面でお見せします。")

# ============================ 7 与える価値 ============================
s = new()
title(s, "tubestudy が与える価値")
vals = [("講師", TINT2, BLUE, "授業へ", ["手作業が減り、授業の準備に時間を回せる", "チェック忘れが、無くなる", "担当が変わっても、記録が残る"]),
        ("生徒", YTINT, INK, "続く", ["出した・見てもらった、が毎回残る", "続けた記録が、続ける力になる", "自分の苦手から、問題が出る"]),
        ("保護者", GTINT, GREEN, "安心", ["見てもらっている、が伝わる", "成績の伸びが、見える", "相談の窓口が、ある"])]
for i, (hd, fill, col, kw, rows) in enumerate(vals):
    x = 0.7 + i * 4.05
    rect(s, x, 1.85, 3.75, 4.6, TINT, radius=0.08)
    pill(s, x + 0.3, 2.15, 1.2, 0.42, hd, 15, fill, col)
    text(s, x + 0.3, 2.85, 3.2, 0.9, [[(kw, 36, True, BLUE)]])
    bullets(s, x + 0.3, 4.05, 3.3, rows, 14, gap=0.18)
notes(s, "価値も、同じ三者で。講師は、手作業が減って授業の準備に時間を回せ、チェック忘れが無くなり、担当が変わっても記録が残る。生徒は、出した・見てもらったが毎回残り、続けた記録が続ける力になり、自分の苦手から問題が出る。保護者は、見てもらっていることが伝わり、成績の伸びが見え、相談の窓口がある。成績が上がり、講師の質が上がり、保護者の評価が上がる。その循環を仕組みで支えたい。")

# ============================ 8 あっと驚く簡単さ ============================
s = new()
title(s, "あっと驚く簡単さ", "覚えることがない。紙のクラスと、そのまま併走できる")
big = [("1回", "評価は、押すだけ", "なぞって、1〜5。\n束の全件に同じ評価が入る。"),
       ("撮って送る", "写真の提出", "撮る → たまる → 送る。\n見開きも2枚で。\nパソコンとスマホを\n行き来しても常に最新。"),
       ("コースごと", "導入は選べる", "講師が「使う」クラスを選ぶ。\n紙のままのクラスと\n半年以上混在して動く。")]
for i, (nm, h, d) in enumerate(big):
    x = 0.7 + i * 3.1
    rect(s, x, 2.05, 2.85, 3.7, TINT, radius=0.08)
    text(s, x + 0.2, 2.3, 2.55, 0.8, [[(nm, 30, True, BLUE)]])
    text(s, x + 0.2, 3.3, 2.55, 0.5, [[(h, 17, True, INK)]])
    para(s, x + 0.2, 3.85, 2.55, d, 12, INK2, spacing=1.4)
phone(s, "s-stage.png", 10.35, 1.95, 2.3, crop_bottom=0.25)
notes(s, "簡単さです。評価は押すだけ。写真は撮って送るだけで、パソコンとスマホを行き来しても常に最新です。導入はコースごとに選べるので、紙のままのクラスと併走しながら少しずつ広げられます。")

# ============================ 9 原理的に安全です ============================
s = new()
title(s, "原理的に安全です", "サーバに入れるのは、カリキュラム（単元の並び）だけ")
safe = [("保護者・生徒の個人情報", "入れない", "氏名・住所・電話は、\nサーバのどこにも無い。\n名簿は、塾の手元の紙だけ。"),
        ("早稲田アカデミーの教材データ", "入れない", "問題や本文は置かない。\n「教材名とページ」と\n正誤だけで回る。"),
        ("カリキュラム", "これだけ", "単元の並びと、\n生徒番号ごとの記録。\n校舎ごとに分けて保管。")]
for i, (h, big_, d) in enumerate(safe):
    x = 0.7 + i * 4.05
    rect(s, x, 1.95, 3.75, 3.35, TINT, radius=0.08)
    text(s, x + 0.35, 2.2, 3.1, 0.5, [[(h, 15, True, INK2)]])
    text(s, x + 0.35, 2.7, 3.1, 0.9, [[(big_, 34, True, BLUE if i < 2 else INK)]])
    para(s, x + 0.35, 3.75, 3.1, d, 13, INK2, spacing=1.45)
rect(s, 0.7, 5.7, 11.9, 1.05, TINT2, radius=0.08)
text(s, 1.0, 5.7, 11.3, 1.05, [[("本部のシステムに全くつながなくても、", 17, True, INK)], [("個別指導塾のようなパーソナライズ宿題と、成績管理が成立します。", 17, True, INK)]],
     anchor=MSO_ANCHOR.MIDDLE, spacing=1.3)
notes(s, "安全について、先にお伝えします。保護者や生徒の個人情報は入れません。氏名・住所・電話はサーバのどこにも無く、名簿は塾の手元の紙だけです。早稲田アカデミーの教材データも入れません。教材名とページ、正誤だけで回ります。入れるのはカリキュラム、つまり単元の並びと、生徒番号ごとの記録だけ。本部のシステムに全くつながなくても、個別指導塾のようなパーソナライズ宿題と成績管理が成立します。")

# ============================ 10 最後に ============================
s = new(NAVY)
text(s, 0.9, 1.7, 11.5, 1.6, [[("先生の時間を、授業へ。", 48, True, WHITE), ("もっと。", 48, True, YEL)]])
text(s, 0.92, 3.1, 11.5, 1.0, [[("一人ひとりに向き合う「仕組み」を、ロンドン校から。", 22, False, ICE)]])
rect(s, 0.9, 4.45, 7.9, 1.1, RGBColor(0x1A, 0x4F, 0x86), radius=0.08)
text(s, 1.2, 4.45, 7.4, 1.1, [[("まず、ロンドン校での試験導入をご検討ください。", 20, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
mark(s, 11.15, 4.4, 1.2)
text(s, 0.92, 6.35, 9, 0.5, [[("秋 建太郎　　時期や進め方は、別途ご相談させてください", 12, False, RGBColor(0xA9, 0xBE, 0xDD))]])
notes(s, "最後に。先生の時間を、授業へ。もっと。一人ひとりに向き合う仕組みを、ロンドン校から。まず、私の担当クラスから、試験的に使わせてください。費用も契約もありません。紙のクラスと混ざったままで動き、いつでも紙に戻せます。3か月後に数字でご報告します。ありがとうございました。")

prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
